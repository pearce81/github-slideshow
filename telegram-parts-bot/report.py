"""Print the reports on your computer, without the bot.

    py report.py                         # top parts for 24h, 7 days, 30 days, this month
    py report.py top 7d                  # most asked in the last 7 days (top 50)
    py report.py top month 100           # top 100 this month
    py report.py count today             # number of messages
    py report.py part 86511-2S000 30d    # how often one part was asked
    py report.py export month            # save a CSV file (opens in Excel)
    py report.py monthly 2025-11 2026-10 100   # top 100 for each month -> Excel file
    py report.py average 2025-11 2026-10 400   # top 400 by average asks per month -> Excel file
    py report.py avg 264102G000 23040-2C250    # average asks per month for these parts (all data)
    py report.py prefix 21910                  # all parts starting with 21910, most asked first
    py report.py prefix 21910 30d              # ...only counting the last 30 days

Periods: today | yesterday | 24h | 7d | 30d | month | lastmonth | all | <N>d |
YYYY-MM-DD [YYYY-MM-DD]
"""

import csv
import os
import sys
from datetime import date, datetime, time
from zoneinfo import ZoneInfo

from parts import extract_part_numbers, normalize_text
from periods import PeriodError, parse_period
from storage import Storage

TZ = ZoneInfo(os.environ.get("TIMEZONE", "UTC"))
TOP_N = int(os.environ.get("TOP_N", "50"))  # how many part numbers each list shows


def pick_chat(db: Storage) -> int:
    if os.environ.get("GROUP_ID"):
        return int(os.environ["GROUP_ID"])
    chats = db.chats()
    if not chats:
        raise SystemExit("The database is empty. Run account_reader.py (or the bot) first.")
    if len(chats) > 1:
        listing = "\n".join(f"  {cid}  ({n} messages)" for cid, n in chats)
        raise SystemExit(f"Several groups in the database; set GROUP_ID to one of:\n{listing}")
    return chats[0][0]


def print_top(rows, label):
    print(f"\nMost asked – {label}")
    if not rows:
        print("  No part numbers found.")
    for i, (part, n, users) in enumerate(rows, 1):
        print(f"  {i:>2}. {part:<14} {n:>4}x  ({users} users)")


def month_starts(first: str, last: str) -> list[date]:
    """['2025-11', '2026-10'] -> 1st day of every month from first to last, inclusive."""
    try:
        y, m = map(int, first.split("-")[:2])
        end_y, end_m = map(int, last.split("-")[:2])
    except ValueError:
        raise PeriodError("Months must look like 2025-11")
    months = []
    while (y, m) <= (end_y, end_m):
        months.append(date(y, m, 1))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    if not months:
        raise PeriodError("The last month is before the first month")
    return months


def monthly_report(db: Storage, chat_id: int, months: list[date], limit: int, folder: str = ".") -> str:
    """Excel file: a summary sheet (parts x months) and one sheet per month. Returns its path."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    bold = Font(bold=True)
    wb = Workbook()
    summary = wb.active
    summary.title = "All months"
    per_month = {}  # month label -> {part: count} for every part asked that month
    top_parts = set()  # parts that made at least one month's top list
    for first in months:
        start, end = month_range(first)
        rows = db.top_parts(chat_id, start, end, limit)
        stats = db.summary(chat_id, start, end)
        label = f"{first:%Y-%m}"
        per_month[label] = {part: n for part, n, _ in db.top_parts(chat_id, start, end, -1)}
        top_parts.update(part for part, _, _ in rows)
        print(f"{first:%B %Y}: {stats['messages']} messages, {stats['part_mentions']} part requests")

        ws = wb.create_sheet(label)
        ws.append([f"Top {limit} part numbers – {first:%B %Y}"])
        ws.append([f"{stats['messages']} messages, {stats['part_mentions']} part requests, "
                   f"{stats['unique_parts']} different part numbers"])
        ws.append([])
        ws.append(["Rank", "Part number", "Times asked", "Different users"])
        for i, (part, n, users) in enumerate(rows, 1):
            ws.append([i, part, n, users])
        ws["A1"].font = bold
        for c in ws[4]:
            c.font = bold
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 13
        ws.column_dimensions["D"].width = 16
        ws.freeze_panes = "A5"

    # Summary: every part that made any month's top list, with its real count in every month.
    labels = list(per_month)
    table = {part: [per_month[m].get(part, 0) for m in labels] for part in top_parts}
    summary.append(["Part number", "Total asked", "Months asked", *labels])
    for part, counts in sorted(table.items(), key=lambda kv: (-sum(kv[1]), kv[0])):
        summary.append([part, sum(counts), sum(1 for n in counts if n), *[n or "" for n in counts]])
    for c in summary[1]:
        c.font = bold
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 12
    summary.column_dimensions["C"].width = 14
    summary.freeze_panes = "B2"

    name = os.path.join(folder, f"monthly_top{limit}_{labels[0]}_to_{labels[-1]}.xlsx")
    wb.save(name)
    return name


def months_and_limit(args: list[str], now: datetime) -> tuple[list[date], int]:
    """[first YYYY-MM] [last YYYY-MM] [N] -> months and N; default: last 12 months including this one."""
    months_args = [a for a in args if "-" in a]
    limit = next((int(a) for a in args if a.isdigit()), TOP_N)
    last = months_args[1] if len(months_args) > 1 else f"{now:%Y-%m}"
    if months_args:
        first = months_args[0]
    else:
        y, m = now.year, now.month - 11
        first = f"{y - 1}-{m + 12}" if m < 1 else f"{y}-{m}"
    return month_starts(first, last), limit


def month_range(first: date):
    nxt = date(first.year + first.month // 12, first.month % 12 + 1, 1)
    return datetime.combine(first, time.min, tzinfo=TZ), datetime.combine(nxt, time.min, tzinfo=TZ)


def classify_months(db: Storage, chat_id: int, months: list[date]):
    """Split months into full ones and partial ones (recording gaps).

    Returns (messages per month, part counts per month, full months, partial months).
    Months with no messages, or under half the messages of a typical (median) month,
    are gaps in the recording and shouldn't count towards an average.
    """
    counts, messages = {}, {}
    for first in months:
        start, end = month_range(first)
        messages[first] = db.summary(chat_id, start, end)["messages"]
        counts[first] = {p.replace("-", ""): n for p, n, _ in db.top_parts(chat_id, start, end, -1)}
    with_data = sorted(m for m in months if messages[m])
    if not with_data:
        raise SystemExit("No messages in these months.")
    typical = sorted(messages[m] for m in with_data)[len(with_data) // 2]
    full = [m for m in with_data if messages[m] >= typical / 2]
    partial = [m for m in with_data if m not in full]
    return messages, counts, full, partial


def part_key(code: str) -> str:
    """Part number as typed (86511-2S000, 865112s000, ۸۶۵۱۱۲S۰۰۰) -> the form without dashes."""
    return normalize_text(code).strip().replace("-", "").replace(" ", "")


def part_averages(db: Storage, chat_id: int, codes: list[str]) -> str:
    """Each part's average asks per month over all full months of data, and its rank, as text."""
    oldest, newest = db.date_range(chat_id)
    if oldest is None:
        raise SystemExit("The database is empty.")
    first, last = datetime.fromtimestamp(oldest, TZ), datetime.fromtimestamp(newest, TZ)
    months = month_starts(f"{first:%Y-%m}", f"{last:%Y-%m}")
    messages, counts, full, partial = classify_months(db, chat_id, months)

    totals = {}
    for m in full:
        for p, n in counts[m].items():
            totals[p] = totals.get(p, 0) + n
    ranking = {p: i for i, (p, _) in enumerate(sorted(totals.items(), key=lambda kv: (-kv[1], kv[0])), 1)}

    out = [f"All data: {first:%Y-%m-%d} to {last:%Y-%m-%d}. Average over {len(full)} full months: "
           + ", ".join(f"{m:%b %Y}" for m in full)]
    if partial:
        out.append("Partial months, not in the average: " + ", ".join(f"{m:%b %Y}" for m in partial))
    for code in codes:
        key = part_key(code)
        per_full = [counts[m].get(key, 0) for m in full]
        total_all = sum(counts[m].get(key, 0) for m in months)
        out.append(f"\n{key}")
        if not total_all:
            out.append("  Never asked in the data.")
            continue
        out.append(f"  Average per month: {sum(per_full) / len(full):.1f}")
        if key in ranking:
            out.append(f"  Rank by average:   {ranking[key]} of {len(ranking)} part numbers")
        out.append(f"  Total asked:       {total_all} (all data)")
        out.append("  Per month:         " + "  ".join(
            f"{m:%b %Y}: {counts[m].get(key, 0)}{' (partial)' if m in partial else ''}"
            for m in months if messages[m]))
    return "\n".join(out)


def prefix_report(db: Storage, chat_id: int, prefix: str, period: list[str], now: datetime) -> str:
    """Every part number starting with prefix, most asked first, as text.

    With no period: all data, plus the average per month over the full months.
    """
    key = part_key(prefix)
    if period:
        start, end, label, _ = parse_period(period, now)
        avg = None
    else:
        oldest, newest = db.date_range(chat_id)
        if oldest is None:
            raise SystemExit("The database is empty.")
        start = datetime.fromtimestamp(oldest, TZ)
        end = datetime.fromtimestamp(newest + 1, TZ)
        label = "all data"
        months = month_starts(f"{start:%Y-%m}", f"{end:%Y-%m}")
        _, counts, full, _ = classify_months(db, chat_id, months)
        avg = lambda p: sum(counts[m].get(p, 0) for m in full) / len(full)  # noqa: E731

    rows = [(p.replace("-", ""), n, users) for p, n, users in db.top_parts(chat_id, start, end, -1)
            if p.replace("-", "").startswith(key)]
    out = [f"Part numbers starting with {key} – {label}: {len(rows)} found, "
           f"{sum(n for _, n, _ in rows)} asks in total"]
    if not rows:
        return out[0]
    if avg:
        out.append(f"(average per month over {len(full)} full months: "
                   + ", ".join(f"{m:%b %Y}" for m in full) + ")")
    out.append(f"\n  {'Rank':>4}  {'Part number':<12} {'Asked':>7} {'Users':>6}" + ("  Avg/month" if avg else ""))
    for i, (p, n, users) in enumerate(rows, 1):
        out.append(f"  {i:>4}  {p:<12} {n:>7} {users:>6}" + (f"  {avg(p):>9.1f}" if avg else ""))
    return "\n".join(out)


def average_report(db: Storage, chat_id: int, months: list[date], limit: int, folder: str = ".") -> str:
    """Excel file ranking parts by average asks per month, over the months with full data.

    Partial months (see classify_months) are shown in grey but not averaged.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    messages, counts, full, partial = classify_months(db, chat_id, months)
    for m in months:
        kind = "used" if m in full else ("partial, not used" if m in partial else "no data, not used")
        print(f"{m:%B %Y}: {messages[m]} messages ({kind})")

    parts = {p for m in full for p in counts[m]}
    table = []
    for p in parts:
        per = [counts[m].get(p, 0) for m in full]
        table.append((p, sum(per) / len(full), sum(per), sum(1 for n in per if n), per,
                      [counts[m].get(p, 0) for m in partial]))
    table.sort(key=lambda r: (-r[1], r[0]))

    wb = Workbook()
    ws = wb.active
    ws.title = "By monthly average"
    ws.append(["Rank", "Part number", "Average per month", f"Total ({len(full)} full months)",
               f"Months asked (of {len(full)})", *[f"{m:%b %Y}" for m in full],
               *[f"{m:%b %Y} (partial)" for m in partial]])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    grey = PatternFill("solid", fgColor="EEEEEE")
    first_partial_col = 6 + len(full)
    for i, (p, avg, total, n_months, per, per_partial) in enumerate(table[:limit], 1):
        ws.append([i, p, round(avg, 1), total, n_months, *per, *per_partial])
        ws.cell(ws.max_row, 2).number_format = "@"  # text: keeps leading zeros
        ws.cell(ws.max_row, 3).number_format = "0.0"
        for col in range(first_partial_col, first_partial_col + len(partial)):
            ws.cell(ws.max_row, col).fill = grey
    ws.column_dimensions["B"].width = 14
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    notes = wb.create_sheet("How it was calculated")
    notes.append([f"Average per month = asks in the {len(full)} full months / {len(full)}."])
    notes.append(["Full months: " + ", ".join(f"{m:%b %Y}" for m in full)])
    notes.append(["Partial months (under half the messages of a typical month), not in the average: "
                  + (", ".join(f"{m:%b %Y} ({messages[m]} messages)" for m in partial) or "none")])
    notes.append(["Months with no data: " + (", ".join(f"{m:%b %Y}" for m in months if not messages[m]) or "none")])
    notes.column_dimensions["A"].width = 120

    name = os.path.join(folder, f"average_top{limit}_{months[0]:%Y-%m}_to_{months[-1]:%Y-%m}.xlsx")
    wb.save(name)
    return name


def main(argv: list[str]) -> None:
    db = Storage(os.environ.get("DB_PATH", "parts.db"))
    chat_id = pick_chat(db)
    now = datetime.now(TZ)
    cmd, args = (argv[0].lower(), argv[1:]) if argv else ("report", [])

    if cmd == "report":
        for key in ("24h", "7d", "30d", "month"):
            start, end, label, _ = parse_period([key], now)
            print_top(db.top_parts(chat_id, start, end, TOP_N), label)
            s = db.summary(chat_id, start, end)
            print(f"  ({s['messages']} messages, {s['part_mentions']} part requests)")
    elif cmd == "top":
        start, end, label, rest = parse_period(args, now)
        limit = int(rest[0]) if rest and rest[0].isdigit() else TOP_N
        print_top(db.top_parts(chat_id, start, end, limit), label)
    elif cmd == "count":
        start, end, label, _ = parse_period(args, now)
        s = db.summary(chat_id, start, end)
        print(f"Group activity – {label}")
        print(f"  Messages:                   {s['messages']}")
        print(f"  Active users:               {s['users']}")
        print(f"  Messages with part numbers: {s['messages_with_parts']}")
        print(f"  Part requests:              {s['part_mentions']}")
        print(f"  Different part numbers:     {s['unique_parts']}")
    elif cmd == "part" and args:
        found = extract_part_numbers(args[0])
        part = found[0] if found else normalize_text(args[0])
        start, end, label, _ = parse_period(args[1:] or ["all"], now)
        n, users = db.part_count(chat_id, part, start, end)
        print(f"{part} was asked {n} times by {users} users ({label}).")
    elif cmd == "export":
        start, end, label, _ = parse_period(args or ["all"], now)
        name = f"part_requests_{start:%Y%m%d}_{end:%Y%m%d}.csv"
        with open(name, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["date", "part_number", "user_id", "user_name", "message_id"])
            for ts, part, uid, uname, mid in db.all_mentions(chat_id, start, end):
                w.writerow([datetime.fromtimestamp(ts, TZ).isoformat(), part, uid, uname, mid])
        print(f"Saved {name} ({label})")
    elif cmd in ("monthly", "average"):
        build = monthly_report if cmd == "monthly" else average_report
        name = build(db, chat_id, *months_and_limit(args, now))
        print(f"\nSaved {os.path.basename(name)}")
        if os.name == "nt":
            os.startfile(name)
    elif cmd == "prefix" and args:
        print(prefix_report(db, chat_id, args[0], args[1:], now))
    elif cmd == "avg" and args:
        print(part_averages(db, chat_id, args))
    else:
        print(__doc__)


if __name__ == "__main__":
    try:
        main(sys.argv[1:])
    except PeriodError as e:
        raise SystemExit(str(e))
